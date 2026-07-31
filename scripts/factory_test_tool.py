# -*- coding: utf-8 -*-
"""
Assembly factory BLE test tool for W2M laser toy devices.

Dependencies:
    pip install PyQt5 bleak openpyxl

This tool intentionally does not auto-connect or auto-reconnect. The operator
scans, selects a device, connects manually, then sends the test or reboot
command.
"""

import asyncio
import csv
import datetime as _dt
import os
import sys
import threading
import uuid
from dataclasses import dataclass
from typing import Dict, Optional, Set

from PyQt5 import QtCore, QtGui, QtWidgets
from PyQt5.QtCore import Qt

try:
    from bleak import BleakClient, BleakScanner
except Exception as ex:  # pragma: no cover - shown in UI at runtime
    BleakClient = None
    BleakScanner = None
    BLEAK_IMPORT_ERROR = ex
else:
    BLEAK_IMPORT_ERROR = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except Exception as ex:  # pragma: no cover - shown in UI at runtime
    Workbook = None
    OPENPYXL_IMPORT_ERROR = ex
else:
    OPENPYXL_IMPORT_ERROR = None


CTRL_PROTO_VERSION = 0x01
CTRL_MSG_TYPE_CMD = 0x01
CTRL_MSG_TYPE_RSP = 0x02
CTRL_MSG_TYPE_EVENT = 0x03

CTRL_STATUS_TEXT = {
    0x00: "OK",
    0x01: "LEN_ERROR",
    0x02: "UNSUPPORTED_CMD",
    0x03: "PARAM_ERROR",
    0x04: "INTERNAL_ERROR",
    0x05: "REJECT_ERROR",
}

CTRL_CMD_DEVICE_REBOOT = 0x5A
CTRL_CMD_FACTORY_TEST_ENTER = 0x5D
CTRL_FACTORY_TEST_MODULE_RADAR = 0x01
CTRL_FACTORY_TEST_MODULE_MOTOR = 0x02
CTRL_FACTORY_TEST_MODULE_LASER = 0x03

RESULT_HEADERS = [
    "时间",
    "设备名称",
    "MAC/地址",
    "MANUFACTURER_DATA",
    "雷达测试",
    "电机测试",
    "激光灯测试",
    "结果",
]

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_RESULT_CSV = os.path.join(SCRIPT_DIR, "factory_test_results.csv")
DEFAULT_RESULT_XLSX = os.path.join(SCRIPT_DIR, "factory_test_results.xlsx")

CTRL_RX_RAW_BYTES = bytes(
    [0x01, 0xA0, 0x0D, 0x0C, 0x0B, 0x0A, 0x09, 0x08, 0x07, 0x06, 0x05, 0x04, 0x03, 0x02, 0x01, 0x00]
)
CTRL_TX_RAW_BYTES = bytes(
    [0x02, 0xA0, 0x0D, 0x0C, 0x0B, 0x0A, 0x09, 0x08, 0x07, 0x06, 0x05, 0x04, 0x03, 0x02, 0x01, 0x00]
)
CTRL_LOG_RAW_BYTES = bytes(
    [0x03, 0xA0, 0x0D, 0x0C, 0x0B, 0x0A, 0x09, 0x08, 0x07, 0x06, 0x05, 0x04, 0x03, 0x02, 0x01, 0x00]
)


def _uuid_candidates(raw: bytes) -> Set[int]:
    return {uuid.UUID(bytes=raw).int, uuid.UUID(bytes=raw[::-1]).int}


def _char_uuid_int(char) -> Optional[int]:
    try:
        return uuid.UUID(str(char.uuid)).int
    except Exception:
        return None


def _find_characteristic(client, raw_uuid: bytes):
    targets = _uuid_candidates(raw_uuid)
    for service in client.services:
        for char in service.characteristics:
            if _char_uuid_int(char) in targets:
                return char
    return None


def _format_manufacturer_data(mfr: Dict[int, bytes]) -> str:
    if not mfr:
        return ""
    parts = []
    for company_id, data in sorted(mfr.items()):
        # Bleak exposes the first two manufacturer bytes as a little-endian
        # company_id. Rebuild the original advertising payload byte order.
        raw = int(company_id).to_bytes(2, "little") + bytes(data)
        parts.append("0x" + raw.hex().upper())
    return "; ".join(parts)


def _build_ctrl_cmd(cmd_id: int, seq: int, payload: bytes = b"") -> bytes:
    plen = len(payload)
    return bytes(
        [
            CTRL_PROTO_VERSION,
            CTRL_MSG_TYPE_CMD,
            cmd_id & 0xFF,
            seq & 0xFF,
            plen & 0xFF,
            (plen >> 8) & 0xFF,
        ]
    ) + payload


def _decode_ctrl_frame(data: bytes) -> Dict[str, object]:
    if len(data) < 6:
        return {
            "short": True,
            "raw": data,
            "text": f"CTRL short: {data.hex(' ').upper()}",
        }
    version, msg_type, cmd_id, seq = data[0], data[1], data[2], data[3]
    payload_len = data[4] | (data[5] << 8)
    payload = data[6 : 6 + payload_len]
    type_name = {CTRL_MSG_TYPE_CMD: "CMD", CTRL_MSG_TYPE_RSP: "RSP", CTRL_MSG_TYPE_EVENT: "EVT"}.get(msg_type, f"0x{msg_type:02X}")
    status = None
    if msg_type == CTRL_MSG_TYPE_RSP and payload:
        status = CTRL_STATUS_TEXT.get(payload[0], f"0x{payload[0]:02X}")
    return {
        "short": False,
        "raw": data,
        "version": version,
        "msg_type": msg_type,
        "cmd_id": cmd_id,
        "seq": seq,
        "payload_len": payload_len,
        "payload": payload,
        "type_name": type_name,
        "status": status,
        "text": f"[{type_name}] cmd=0x{cmd_id:02X} seq={seq} len={payload_len}"
        + (f" status={status}" if status is not None else "")
        + f" payload={payload.hex(' ').upper()}",
    }


@dataclass
class ScanDevice:
    name: str
    address: str
    rssi: Optional[int]
    manufacturer_data: str


class BleWorker(QtCore.QObject):
    devices_changed = QtCore.pyqtSignal(list)
    log_line = QtCore.pyqtSignal(str)
    firmware_log_line = QtCore.pyqtSignal(str)
    status_changed = QtCore.pyqtSignal(str, bool)
    ctrl_frame = QtCore.pyqtSignal(dict)

    def __init__(self):
        super().__init__()
        self._loop = asyncio.new_event_loop()
        self._thread = threading.Thread(target=self._run_loop, daemon=True)
        self._thread.start()
        self._client = None
        self._rx_char = None
        self._seq = 0
        self._log_buf = bytearray()
        self._devices: Dict[str, ScanDevice] = {}

    def _run_loop(self):
        asyncio.set_event_loop(self._loop)
        self._loop.run_forever()

    def stop(self):
        try:
            self.disconnect()
        finally:
            self._loop.call_soon_threadsafe(self._loop.stop)

    def _schedule(self, coro):
        return asyncio.run_coroutine_threadsafe(coro, self._loop)

    def scan(self, timeout_s: float = 1.0):
        self._schedule(self._scan(timeout_s))

    def connect(self, address: str):
        self._schedule(self._connect(address))

    def disconnect(self):
        self._schedule(self._disconnect())

    def send_factory_test_enter(self):
        self._schedule(self._send_cmd(CTRL_CMD_FACTORY_TEST_ENTER, '进入测试模式'))

    def send_factory_test_module(self, module: int, enable: bool):
        labels = {
            CTRL_FACTORY_TEST_MODULE_RADAR: '雷达',
            CTRL_FACTORY_TEST_MODULE_MOTOR: '电机',
            CTRL_FACTORY_TEST_MODULE_LASER: '激光灯',
        }
        label = labels.get(module, f'模块0x{module:02X}')
        suffix = '开启' if enable else '关闭'
        payload = bytes([module & 0xFF, 1 if enable else 0])
        self._schedule(self._send_cmd(CTRL_CMD_FACTORY_TEST_ENTER, f'{label}{suffix}', payload))

    def send_reboot(self):
        self._schedule(self._send_cmd(CTRL_CMD_DEVICE_REBOOT, '重启设备'))

    async def _scan(self, timeout_s: float):
        if BLEAK_IMPORT_ERROR is not None:
            self.log_line.emit(f'Bleak 未安装或导入失败: {BLEAK_IMPORT_ERROR}')
            return

        self.log_line.emit(f'开始扫描 W2MLaserTOY 设备，{timeout_s:.0f}s...')
        found: Dict[str, ScanDevice] = {}

        def on_adv(device, adv):
            name = adv.local_name or getattr(device, 'name', None) or ''
            if not name.startswith('W2MLaserTOY'):
                return
            rssi = getattr(adv, 'rssi', None)
            if rssi is None:
                rssi = getattr(device, 'rssi', None)
            found[device.address] = ScanDevice(
                name=name,
                address=device.address,
                rssi=rssi,
                manufacturer_data=_format_manufacturer_data(getattr(adv, 'manufacturer_data', {}) or {}),
            )
            self._devices = dict(found)
            self.devices_changed.emit(list(found.values()))

        scanner = BleakScanner(detection_callback=on_adv)
        try:
            await scanner.start()
            await asyncio.sleep(timeout_s)
            await scanner.stop()
        except Exception as ex:
            self.log_line.emit(f'扫描失败: {ex}')
            return

        self._devices = dict(found)
        self.devices_changed.emit(list(found.values()))
        self.log_line.emit(f'扫描完成，找到 {len(found)} 个 W2MLaserTOY 设备')

    async def _connect(self, address: str):
        if BLEAK_IMPORT_ERROR is not None:
            self.log_line.emit(f'Bleak 未安装或导入失败: {BLEAK_IMPORT_ERROR}')
            return
        if not address:
            self.log_line.emit('请先选择设备')
            return
        if self._client and self._client.is_connected:
            self.log_line.emit('已连接，请先断开')
            return

        self.log_line.emit(f'连接中: {address}')

        def on_disconnect(_client):
            self._client = None
            self._rx_char = None
            self.status_changed.emit('已断开', False)
            self.log_line.emit('设备已断开')

        client = BleakClient(address, disconnected_callback=on_disconnect)
        try:
            await client.connect()
            get_services = getattr(client, 'get_services', None)
            if get_services:
                out = get_services()
                if asyncio.iscoroutine(out):
                    await out
            else:
                _ = client.services

            self._rx_char = _find_characteristic(client, CTRL_RX_RAW_BYTES)
            tx_char = _find_characteristic(client, CTRL_TX_RAW_BYTES)
            log_char = _find_characteristic(client, CTRL_LOG_RAW_BYTES)
            if self._rx_char is None:
                raise RuntimeError('未找到 Ctrl RX 写特征')
            if tx_char is not None:
                await client.start_notify(tx_char, self._on_ctrl_notify)
            if log_char is not None:
                await client.start_notify(log_char, self._on_log_notify)
            else:
                self.log_line.emit('未找到 BLE_LOG_D 日志特征，仅显示控制响应')

            self._client = client
            self.status_changed.emit(f'已连接 {address}', True)
            self.log_line.emit('连接成功')
        except Exception as ex:
            try:
                await client.disconnect()
            except Exception:
                pass
            self._client = None
            self._rx_char = None
            self.status_changed.emit('连接失败', False)
            self.log_line.emit(f'连接失败: {ex}')

    async def _disconnect(self):
        client = self._client
        self._client = None
        self._rx_char = None
        if client is None:
            self.status_changed.emit('未连接', False)
            return
        try:
            if client.is_connected:
                await client.disconnect()
        except Exception as ex:
            self.log_line.emit(f'断开失败: {ex}')
            return
        self.status_changed.emit('已断开', False)
        self.log_line.emit('已手动断开')

    async def _send_cmd(self, cmd_id: int, label: str, payload: bytes = b''):
        client = self._client
        if client is None or not client.is_connected or self._rx_char is None:
            self.log_line.emit('未连接，无法发送命令')
            return
        seq = self._seq & 0xFF
        self._seq = (self._seq + 1) & 0xFF
        frame = _build_ctrl_cmd(cmd_id, seq, payload)
        try:
            await client.write_gatt_char(self._rx_char, frame, response=False)
            self.log_line.emit(f'已发送 {label}: cmd=0x{cmd_id:02X} seq={seq}')
        except Exception as ex:
            self.log_line.emit(f'发送失败 {label}: {ex}')

    def _on_ctrl_notify(self, _sender, data: bytearray):
        frame = _decode_ctrl_frame(bytes(data))
        self.ctrl_frame.emit(frame)
        self.log_line.emit(frame['text'])

    def _on_log_notify(self, _sender, data: bytearray):
        self._log_buf.extend(bytes(data))
        while b'\n' in self._log_buf:
            idx = self._log_buf.index(0x0A)
            raw = bytes(self._log_buf[:idx]).rstrip(b'\r')
            del self._log_buf[: idx + 1]
            if raw:
                self.firmware_log_line.emit(raw.decode('utf-8', errors='replace'))
        if len(self._log_buf) > 512:
            raw = bytes(self._log_buf)
            self._log_buf.clear()
            self.firmware_log_line.emit(raw.decode('utf-8', errors='replace'))


class FactoryTestWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.worker = BleWorker()
        self.devices: Dict[str, ScanDevice] = {}
        self.connected_address: Optional[str] = None
        self.factory_ready = False
        self.tested_mfr_set = set()
        self.module_results = {
            'radar': None,
            'motor': None,
            'laser': None,
        }
        self.ui_font_size = 20
        self.result_csv_path = DEFAULT_RESULT_CSV
        self.result_xlsx_path = DEFAULT_RESULT_XLSX

        self.setWindowTitle('W2MLaserTOY 组装工厂测试工具')
        self.resize(1280, 860)
        self._build_ui()
        self._connect_signals()
        self._set_connected(False)
        self._update_result_path_label()
        self._refresh_csv_preview()

        if BLEAK_IMPORT_ERROR is not None:
            self._append_log(f'Bleak 未安装或导入失败: {BLEAK_IMPORT_ERROR}')

    def _build_stylesheet(self, font_px: int) -> str:
        return f"""
            QWidget {{
                font-size: {font_px}px;
                color: #e5e7eb;
                background: #0f172a;
            }}
            QGroupBox {{
                font-weight: 600;
                border: 1px solid #334155;
                border-radius: 6px;
                margin-top: 10px;
                background: #111827;
            }}
            QGroupBox::title {{
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 4px;
                color: #cbd5e1;
            }}
            QPushButton {{
                min-width: 76px;
                padding: 6px 10px;
                border: 1px solid #475569;
                border-radius: 5px;
                background: #1f2937;
                color: #e5e7eb;
            }}
            QPushButton:hover {{ background: #334155; }}
            QPushButton:pressed {{ background: #475569; }}
            QPushButton:disabled {{
                color: #64748b;
                background: #1e293b;
                border-color: #334155;
            }}
            QPlainTextEdit, QTableWidget, QTreeWidget, QComboBox, QSpinBox {{
                border: 1px solid #334155;
                border-radius: 4px;
                background: #0b1220;
                color: #e5e7eb;
            }}
            QPlainTextEdit, QTableWidget, QTreeWidget {{
                selection-background-color: #2563eb;
                selection-color: #ffffff;
            }}
            QTreeWidget {{
                alternate-background-color: #111827;
            }}
            QTreeWidget::item {{
                padding: 2px 4px;
                color: #e5e7eb;
            }}
            QTreeWidget::item:alternate {{
                background: #111827;
            }}
            QTreeWidget::item:selected {{
                background: #1d4ed8;
                color: #ffffff;
            }}
            QTreeWidget:disabled {{
                background: #0b1220;
                color: #cbd5e1;
            }}
            QTreeWidget::item:disabled {{
                color: #cbd5e1;
            }}
            QHeaderView::section {{
                background: #1e293b;
                color: #e2e8f0;
                border: 1px solid #334155;
                padding: 6px 8px;
                font-weight: 600;
            }}
            QTableWidget::item:selected {{
                background: #1d4ed8;
                color: #ffffff;
            }}
            QComboBox::drop-down, QSpinBox::up-button, QSpinBox::down-button {{
                border-left: 1px solid #334155;
                width: 20px;
            }}
            QComboBox QAbstractItemView {{
                background: #0b1220;
                color: #e5e7eb;
                selection-background-color: #2563eb;
                border: 1px solid #334155;
            }}
            QLabel#statusLabel {{
                padding: 6px 10px;
                border-radius: 5px;
                background: #1e293b;
                color: #bfdbfe;
                font-weight: 600;
            }}
            QSplitter::handle {{
                background: #334155;
            }}
        """

    def _apply_ui_style(self):
        app = QtWidgets.QApplication.instance()
        if app is not None:
            app_font = app.font()
            app_font.setPointSize(self.ui_font_size)
            app.setFont(app_font)
        self.setStyleSheet(self._build_stylesheet(self.ui_font_size))

    def _on_font_size_changed(self, value: int):
        self.ui_font_size = value
        self._apply_ui_style()
        self.device_tree.resizeColumnToContents(0)
        self.device_tree.resizeColumnToContents(1)
        self.device_tree.resizeColumnToContents(2)
        self.device_tree.resizeColumnToContents(3)
        self.csv_preview_table.resizeRowsToContents()

    def _build_ui(self):
        central = QtWidgets.QWidget()
        self.setCentralWidget(central)
        root = QtWidgets.QVBoxLayout(central)
        root.setContentsMargins(12, 12, 12, 12)
        root.setSpacing(10)
        self._apply_ui_style()
        self.setStyleSheet("""
            QWidget {
                font-size: 13px;
                color: #e5e7eb;
                background: #0f172a;
            }
            QGroupBox {
                font-weight: 600;
                border: 1px solid #334155;
                border-radius: 6px;
                margin-top: 10px;
                background: #111827;
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 4px;
                color: #cbd5e1;
            }
            QPushButton {
                min-width: 76px;
                padding: 6px 10px;
                border: 1px solid #475569;
                border-radius: 5px;
                background: #1f2937;
                color: #e5e7eb;
            }
            QPushButton:hover { background: #334155; }
            QPushButton:pressed { background: #475569; }
            QPushButton:disabled {
                color: #64748b;
                background: #1e293b;
                border-color: #334155;
            }
            QPlainTextEdit, QTableWidget, QTreeWidget, QComboBox {
                border: 1px solid #334155;
                border-radius: 4px;
                background: #0b1220;
                color: #e5e7eb;
            }
            QPlainTextEdit, QTableWidget, QTreeWidget {
                selection-background-color: #2563eb;
                selection-color: #ffffff;
            }
            QHeaderView::section {
                background: #1e293b;
                color: #e2e8f0;
                border: 1px solid #334155;
                padding: 6px 8px;
                font-weight: 600;
            }
            QTableWidget::item:selected, QTreeWidget::item:selected {
                background: #1d4ed8;
                color: #ffffff;
            }
            QComboBox::drop-down {
                border-left: 1px solid #334155;
                width: 20px;
            }
            QComboBox QAbstractItemView {
                background: #0b1220;
                color: #e5e7eb;
                selection-background-color: #2563eb;
                border: 1px solid #334155;
            }
            QLabel#statusLabel {
                padding: 6px 10px;
                border-radius: 5px;
                background: #1e293b;
                color: #bfdbfe;
                font-weight: 600;
            }
            QSplitter::handle {
                background: #334155;
            }
        """)

        device_box = QtWidgets.QGroupBox('设备选择')
        self._apply_ui_style()
        device_layout = QtWidgets.QVBoxLayout(device_box)
        scan_row = QtWidgets.QHBoxLayout()
        self.scan_btn = QtWidgets.QPushButton('扫描 W2MLaserTOY 设备')
        self.device_combo = QtWidgets.QComboBox()
        self.device_combo.setMinimumWidth(380)
        self.connect_btn = QtWidgets.QPushButton('连接')
        self.disconnect_btn = QtWidgets.QPushButton('断开')
        scan_row.addWidget(self.scan_btn)
        scan_row.addWidget(self.device_combo, 1)
        scan_row.addWidget(self.connect_btn)
        scan_row.addWidget(self.disconnect_btn)
        device_layout.addLayout(scan_row)

        self.status_label = QtWidgets.QLabel('未连接')
        self.status_label.setObjectName('statusLabel')
        device_layout.addWidget(self.status_label)
        font_row = QtWidgets.QHBoxLayout()
        font_row.setSpacing(8)
        self.font_size_label = QtWidgets.QLabel('字体大小')
        self.font_size_spin = QtWidgets.QSpinBox()
        self.font_size_spin.setRange(15, 25)
        self.font_size_spin.setValue(self.ui_font_size)
        self.font_size_spin.setSuffix(' px')
        self.font_size_spin.setFixedWidth(96)
        font_row.addWidget(self.font_size_label)
        font_row.addWidget(self.font_size_spin)
        font_row.addStretch(1)
        device_layout.addLayout(font_row)

        self.device_tree = QtWidgets.QTreeWidget()
        self.device_tree.setHeaderLabels(['名称', 'MAC/地址', 'RSSI', '是否测试过', 'MANUFACTURER_DATA'])
        self.device_tree.setRootIsDecorated(False)
        self.device_tree.setAlternatingRowColors(True)
        self.device_tree.setUniformRowHeights(True)
        self.device_tree.header().setSectionResizeMode(0, QtWidgets.QHeaderView.ResizeToContents)
        self.device_tree.header().setSectionResizeMode(1, QtWidgets.QHeaderView.ResizeToContents)
        self.device_tree.header().setSectionResizeMode(2, QtWidgets.QHeaderView.ResizeToContents)
        self.device_tree.header().setSectionResizeMode(3, QtWidgets.QHeaderView.ResizeToContents)
        self.device_tree.header().setSectionResizeMode(4, QtWidgets.QHeaderView.Stretch)
        device_layout.addWidget(self.device_tree, 2)
        root.addWidget(device_box, 3)

        action_box = QtWidgets.QGroupBox('设备操作')
        action_box_layout = QtWidgets.QVBoxLayout(action_box)
        action_row = QtWidgets.QHBoxLayout()
        self.start_test_btn = QtWidgets.QPushButton('开始测试')
        self.reboot_btn = QtWidgets.QPushButton('重启设备')
        self.clear_log_btn = QtWidgets.QPushButton('清空日志')
        self.start_test_btn.setMinimumHeight(40)
        self.reboot_btn.setMinimumHeight(40)
        action_row.addWidget(self.start_test_btn)
        action_row.addWidget(self.reboot_btn)
        action_row.addStretch(1)
        action_row.addWidget(self.clear_log_btn)
        action_box_layout.addLayout(action_row)
        root.addWidget(action_box)

        module_box = QtWidgets.QGroupBox('模块控制')
        test_ops_row = QtWidgets.QHBoxLayout()
        test_ops_row.setSpacing(10)

        module_row = QtWidgets.QVBoxLayout(module_box)
        module_row.setSpacing(10)
        self.radar_on_btn = QtWidgets.QPushButton('雷达开启')
        self.radar_off_btn = QtWidgets.QPushButton('雷达关闭')
        self.motor_on_btn = QtWidgets.QPushButton('电机开启')
        self.motor_off_btn = QtWidgets.QPushButton('电机关闭')
        self.laser_on_btn = QtWidgets.QPushButton('激光灯开启')
        self.laser_off_btn = QtWidgets.QPushButton('激光灯关闭')
        for btn in (self.radar_on_btn, self.radar_off_btn, self.motor_on_btn, self.motor_off_btn, self.laser_on_btn, self.laser_off_btn):
            btn.setMinimumHeight(34)
        self._add_button_pair(module_row, '雷达控制', self.radar_on_btn, self.radar_off_btn)
        self._add_button_pair(module_row, '电机控制', self.motor_on_btn, self.motor_off_btn)
        self._add_button_pair(module_row, '激光灯控制', self.laser_on_btn, self.laser_off_btn)
        test_ops_row.addWidget(module_box, 1)

        result_box = QtWidgets.QGroupBox('测试结果确认')
        result_row = QtWidgets.QVBoxLayout(result_box)
        result_row.setSpacing(10)
        self.radar_ok_btn = QtWidgets.QPushButton('雷达成功')
        self.radar_fail_btn = QtWidgets.QPushButton('雷达失败')
        self.motor_ok_btn = QtWidgets.QPushButton('电机成功')
        self.motor_fail_btn = QtWidgets.QPushButton('电机失败')
        self.laser_ok_btn = QtWidgets.QPushButton('激光灯成功')
        self.laser_fail_btn = QtWidgets.QPushButton('激光灯失败')
        self.confirm_btn = QtWidgets.QPushButton('测试结果保存')
        self.export_excel_btn = QtWidgets.QPushButton('CSV 转 Excel')
        for btn in (self.radar_ok_btn, self.radar_fail_btn, self.motor_ok_btn, self.motor_fail_btn, self.laser_ok_btn, self.laser_fail_btn, self.confirm_btn, self.export_excel_btn):
            btn.setMinimumHeight(34)
        self._add_button_pair(result_row, '雷达结果', self.radar_ok_btn, self.radar_fail_btn)
        self._add_button_pair(result_row, '电机结果', self.motor_ok_btn, self.motor_fail_btn)
        self._add_button_pair(result_row, '激光灯结果', self.laser_ok_btn, self.laser_fail_btn)
        result_action_row = QtWidgets.QHBoxLayout()
        result_action_row.setSpacing(10)
        result_action_row.addWidget(self.confirm_btn)
        result_action_row.addWidget(self.export_excel_btn)
        result_row.addLayout(result_action_row)
        self.result_path_label = QtWidgets.QLabel()
        self.result_path_label.setTextInteractionFlags(Qt.TextSelectableByMouse)
        self.choose_csv_btn = QtWidgets.QPushButton('选择 CSV')
        self.choose_excel_btn = QtWidgets.QPushButton('选择 Excel')
        for btn in (self.choose_csv_btn, self.choose_excel_btn):
            btn.setMinimumHeight(30)
        path_row = QtWidgets.QHBoxLayout()
        path_row.setSpacing(8)
        path_row.addWidget(self.result_path_label, 1)
        path_row.addWidget(self.choose_csv_btn)
        path_row.addWidget(self.choose_excel_btn)
        result_panel = QtWidgets.QWidget()
        result_layout = QtWidgets.QVBoxLayout(result_panel)
        result_layout.setContentsMargins(0, 0, 0, 0)
        result_layout.addWidget(result_box)
        result_layout.addLayout(path_row)
        test_ops_row.addWidget(result_panel, 1)
        root.addLayout(test_ops_row)

        preview_box = QtWidgets.QGroupBox('CSV 实时预览')
        preview_layout = QtWidgets.QVBoxLayout(preview_box)
        self.csv_preview_table = QtWidgets.QTableWidget()
        self.csv_preview_table.setColumnCount(len(RESULT_HEADERS))
        self.csv_preview_table.setHorizontalHeaderLabels(RESULT_HEADERS)
        self.csv_preview_table.horizontalHeader().setSectionResizeMode(QtWidgets.QHeaderView.Stretch)
        self.csv_preview_table.setAlternatingRowColors(True)
        self.csv_preview_table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.csv_preview_table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        preview_layout.addWidget(self.csv_preview_table)
        root.addWidget(preview_box, 2)

        log_splitter = QtWidgets.QSplitter(Qt.Horizontal)
        host_log_box = QtWidgets.QGroupBox('上位机日志')
        host_log_layout = QtWidgets.QVBoxLayout(host_log_box)
        self.host_log_text = QtWidgets.QPlainTextEdit()
        self.host_log_text.setReadOnly(True)
        self.host_log_text.setPlaceholderText('扫描、连接、按钮操作和控制响应会显示在这里')
        host_log_layout.addWidget(self.host_log_text)

        firmware_log_box = QtWidgets.QGroupBox('固件上传日志')
        firmware_log_layout = QtWidgets.QVBoxLayout(firmware_log_box)
        self.firmware_log_text = QtWidgets.QPlainTextEdit()
        self.firmware_log_text.setReadOnly(True)
        self.firmware_log_text.setPlaceholderText('BLE_LOG_D 固件日志会显示在这里')
        firmware_log_layout.addWidget(self.firmware_log_text)

        log_splitter.addWidget(host_log_box)
        log_splitter.addWidget(firmware_log_box)
        log_splitter.setStretchFactor(0, 1)
        log_splitter.setStretchFactor(1, 1)
        root.addWidget(log_splitter, 3)

    def _add_button_pair(self, parent_layout, title: str, left_btn, right_btn):
        box = QtWidgets.QGroupBox(title)
        layout = QtWidgets.QHBoxLayout(box)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(6)
        layout.addWidget(left_btn)
        layout.addWidget(right_btn)
        parent_layout.addWidget(box, 1)

    def _connect_signals(self):
        self.scan_btn.clicked.connect(lambda: self.worker.scan())
        self.connect_btn.clicked.connect(self._on_connect_clicked)
        self.disconnect_btn.clicked.connect(self.worker.disconnect)
        self.start_test_btn.clicked.connect(self._on_start_test_clicked)
        self.reboot_btn.clicked.connect(self.worker.send_reboot)
        self.clear_log_btn.clicked.connect(self._clear_logs)
        self.device_tree.itemDoubleClicked.connect(self._on_device_double_clicked)
        self.font_size_spin.valueChanged.connect(self._on_font_size_changed)

        self.radar_on_btn.clicked.connect(lambda: self._send_module_ctrl(CTRL_FACTORY_TEST_MODULE_RADAR, True))
        self.radar_off_btn.clicked.connect(lambda: self._send_module_ctrl(CTRL_FACTORY_TEST_MODULE_RADAR, False))
        self.motor_on_btn.clicked.connect(lambda: self._send_module_ctrl(CTRL_FACTORY_TEST_MODULE_MOTOR, True))
        self.motor_off_btn.clicked.connect(lambda: self._send_module_ctrl(CTRL_FACTORY_TEST_MODULE_MOTOR, False))
        self.laser_on_btn.clicked.connect(lambda: self._send_module_ctrl(CTRL_FACTORY_TEST_MODULE_LASER, True))
        self.laser_off_btn.clicked.connect(lambda: self._send_module_ctrl(CTRL_FACTORY_TEST_MODULE_LASER, False))

        self.radar_ok_btn.clicked.connect(lambda: self._set_module_result('radar', '成功'))
        self.radar_fail_btn.clicked.connect(lambda: self._set_module_result('radar', '失败'))
        self.motor_ok_btn.clicked.connect(lambda: self._set_module_result('motor', '成功'))
        self.motor_fail_btn.clicked.connect(lambda: self._set_module_result('motor', '失败'))
        self.laser_ok_btn.clicked.connect(lambda: self._set_module_result('laser', '成功'))
        self.laser_fail_btn.clicked.connect(lambda: self._set_module_result('laser', '失败'))
        self.confirm_btn.clicked.connect(self._on_confirm_clicked)
        self.export_excel_btn.clicked.connect(self._on_export_excel_clicked)
        self.choose_csv_btn.clicked.connect(self._on_choose_csv_clicked)
        self.choose_excel_btn.clicked.connect(self._on_choose_excel_clicked)

        self.worker.devices_changed.connect(self._on_devices_changed)
        self.worker.log_line.connect(self._append_log)
        self.worker.firmware_log_line.connect(self._append_firmware_log)
        self.worker.status_changed.connect(self._on_status_changed)
        self.worker.ctrl_frame.connect(self._on_ctrl_frame)

    def _on_connect_clicked(self):
        address = self.device_combo.currentData()
        self.worker.connect(address or '')

    def _on_device_double_clicked(self, item, _column):
        address = item.data(0, Qt.UserRole)
        idx = self.device_combo.findData(address)
        if idx >= 0:
            self.device_combo.setCurrentIndex(idx)

    def _on_start_test_clicked(self):
        self.factory_ready = False
        self._reset_module_results()
        self.worker.send_factory_test_enter()
        self._append_log('已发送进入测试模式命令，等待设备响应')
        self._update_buttons()

    def _send_module_ctrl(self, module: int, enable: bool):
        if not self.factory_ready:
            self._append_log('设备尚未进入测试状态，不能操作模块')
            return
        self.worker.send_factory_test_module(module, enable)

    def _module_name(self, module: str) -> str:
        return {'radar': '雷达', 'motor': '电机', 'laser': '激光灯'}[module]

    def _on_ctrl_frame(self, frame: dict):
        if frame.get('short'):
            return
        if frame.get('msg_type') != CTRL_MSG_TYPE_RSP or frame.get('cmd_id') != CTRL_CMD_FACTORY_TEST_ENTER:
            return
        payload = frame.get('payload', b'') or b''
        status = frame.get('status')
        if len(payload) <= 1:
            if status == 'OK':
                self.factory_ready = True
                self._append_log('设备已进入测试状态')
            else:
                self.factory_ready = False
                self._append_log(f'进入测试状态失败: {status}')
        elif len(payload) >= 3 and status == 'OK':
            module = payload[1]
            enable = bool(payload[2])
            module_name = {
                CTRL_FACTORY_TEST_MODULE_RADAR: '雷达',
                CTRL_FACTORY_TEST_MODULE_MOTOR: '电机',
                CTRL_FACTORY_TEST_MODULE_LASER: '激光灯',
            }.get(module, f'模块0x{module:02X}')
            self._append_log(f'{module_name}{"开启" if enable else "关闭"}已执行')
        self._update_buttons()

    def _on_devices_changed(self, devices):
        current = self.device_combo.currentData()
        self.devices = {d.address: d for d in devices}
        self._refresh_tested_mfr_set()

        self.device_combo.blockSignals(True)
        self.device_combo.clear()
        for dev in devices:
            self.device_combo.addItem(f'{dev.name}  {dev.address}', dev.address)
        if current:
            idx = self.device_combo.findData(current)
            if idx >= 0:
                self.device_combo.setCurrentIndex(idx)
        self.device_combo.blockSignals(False)

        self.device_tree.clear()
        for row_idx, dev in enumerate(devices):
            tested = '已测' if dev.manufacturer_data and dev.manufacturer_data in self.tested_mfr_set else '未测'
            item = QtWidgets.QTreeWidgetItem([
                dev.name,
                dev.address,
                '' if dev.rssi is None else str(dev.rssi),
                tested,
                dev.manufacturer_data,
            ])
            row_bg = QtGui.QBrush(QtGui.QColor('#0b1220' if row_idx % 2 == 0 else '#111827'))
            row_fg = QtGui.QBrush(QtGui.QColor('#e5e7eb'))
            for col_idx in range(5):
                item.setBackground(col_idx, row_bg)
                item.setForeground(col_idx, row_fg)
            if tested == '已测':
                item.setForeground(3, QtGui.QBrush(QtGui.QColor('#86efac')))
            else:
                item.setForeground(3, QtGui.QBrush(QtGui.QColor('#fbbf24')))
            item.setData(0, Qt.UserRole, dev.address)
            self.device_tree.addTopLevelItem(item)
        if self.device_combo.count() > 0 and self.device_combo.currentIndex() < 0:
            self.device_combo.setCurrentIndex(0)

    def _on_status_changed(self, text: str, connected: bool):
        self.status_label.setText(text)
        if connected:
            self.connected_address = self.device_combo.currentData()
        else:
            self.connected_address = None
            self.factory_ready = False
            self._reset_module_results()
        self._set_connected(connected)

    def _set_connected(self, connected: bool):
        self.scan_btn.setEnabled(not connected)
        self.device_combo.setEnabled(not connected)
        self.device_tree.setEnabled(not connected)
        self.connect_btn.setEnabled(not connected)
        self.disconnect_btn.setEnabled(connected)
        self.start_test_btn.setEnabled(connected)
        self.reboot_btn.setEnabled(connected)
        self._update_buttons()

    def _reset_module_results(self):
        self.module_results = {'radar': None, 'motor': None, 'laser': None}

    def _set_module_result(self, module: str, result: str):
        if not self.factory_ready:
            self._append_log('设备尚未进入测试状态，不能记录结果')
            return
        self.module_results[module] = result
        self._append_log(f'{self._module_name(module)}测试结果：{result}')
        self._update_buttons()

    def _update_buttons(self):
        can_control = self.connected_address is not None and self.factory_ready
        for btn in (
            self.radar_on_btn, self.radar_off_btn,
            self.motor_on_btn, self.motor_off_btn,
            self.laser_on_btn, self.laser_off_btn,
            self.radar_ok_btn, self.radar_fail_btn,
            self.motor_ok_btn, self.motor_fail_btn,
            self.laser_ok_btn, self.laser_fail_btn,
        ):
            btn.setEnabled(can_control)
        self.confirm_btn.setEnabled(
            can_control and all(self.module_results[k] in ('成功', '失败') for k in ('radar', 'motor', 'laser'))
        )

    def _current_device(self) -> Optional[ScanDevice]:
        if not self.connected_address:
            return None
        return self.devices.get(self.connected_address) or ScanDevice(name='', address=self.connected_address, rssi=None, manufacturer_data='')

    def _build_result_row(self, dev: ScanDevice):
        return [
            _dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            dev.name,
            dev.address,
            dev.manufacturer_data,
            self.module_results['radar'] or '',
            self.module_results['motor'] or '',
            self.module_results['laser'] or '',
            'PASS' if all(self.module_results[k] == '成功' for k in ('radar', 'motor', 'laser')) else 'FAIL',
        ]

    def _short_path(self, path: str) -> str:
        base = os.path.basename(path)
        folder = os.path.basename(os.path.dirname(path))
        if folder:
            return os.path.join(folder, base)
        return base or path

    def _update_result_path_label(self):
        csv_short = self._short_path(self.result_csv_path)
        excel_short = self._short_path(self.result_xlsx_path)
        self.result_path_label.setText(f'CSV: {csv_short}    Excel: {excel_short}')
        self.result_path_label.setToolTip(f'CSV: {self.result_csv_path}\nExcel: {self.result_xlsx_path}')

    def _normalize_save_path(self, path: str, suffix: str) -> str:
        if not path:
            return path
        root, ext = os.path.splitext(path)
        if not ext:
            return root + suffix
        return path

    def _on_choose_csv_clicked(self):
        path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self,
            '选择 CSV 保存路径',
            self.result_csv_path,
            'CSV Files (*.csv);;All Files (*)',
        )
        path = self._normalize_save_path(path, '.csv')
        if not path:
            return
        self.result_csv_path = path
        self._update_result_path_label()
        self._refresh_csv_preview()
        self._on_devices_changed(list(self.devices.values()))

    def _on_choose_excel_clicked(self):
        path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self,
            '选择 Excel 保存路径',
            self.result_xlsx_path,
            'Excel Files (*.xlsx);;All Files (*)',
        )
        path = self._normalize_save_path(path, '.xlsx')
        if not path:
            return
        self.result_xlsx_path = path
        self._update_result_path_label()

    def _read_csv_rows(self):
        if not os.path.exists(self.result_csv_path):
            return []
        with open(self.result_csv_path, 'r', newline='', encoding='utf-8-sig') as fp:
            rows = list(csv.reader(fp))
        if not rows:
            return []
        if rows[0] == RESULT_HEADERS:
            return rows[1:]
        return rows

    def _write_csv_rows(self, rows):
        os.makedirs(os.path.dirname(os.path.abspath(self.result_csv_path)), exist_ok=True)
        with open(self.result_csv_path, 'w', newline='', encoding='utf-8-sig') as fp:
            writer = csv.writer(fp)
            writer.writerow(RESULT_HEADERS)
            writer.writerows(rows)

    def _refresh_tested_mfr_set(self, rows=None):
        if rows is None:
            rows = self._read_csv_rows()
        mfr_col = RESULT_HEADERS.index('MANUFACTURER_DATA')
        self.tested_mfr_set = {
            row[mfr_col]
            for row in rows
            if len(row) > mfr_col and row[mfr_col]
        }

    def _save_csv_record(self, dev: ScanDevice):
        rows = self._read_csv_rows()
        if dev.manufacturer_data:
            mfr_col = RESULT_HEADERS.index('MANUFACTURER_DATA')
            rows = [row for row in rows if len(row) <= mfr_col or row[mfr_col] != dev.manufacturer_data]
        rows.append(self._build_result_row(dev))
        self._write_csv_rows(rows)
        self._refresh_csv_preview(rows)
        self._refresh_tested_mfr_set(rows)
        self._on_devices_changed(list(self.devices.values()))

    def _refresh_csv_preview(self, rows=None):
        if rows is None:
            rows = self._read_csv_rows()
        self._refresh_tested_mfr_set(rows)
        self.csv_preview_table.setRowCount(len(rows))
        for row_idx, row in enumerate(rows):
            normalized = list(row[:len(RESULT_HEADERS)]) + [''] * max(0, len(RESULT_HEADERS) - len(row))
            for col_idx, value in enumerate(normalized[:len(RESULT_HEADERS)]):
                item = QtWidgets.QTableWidgetItem(str(value))
                self.csv_preview_table.setItem(row_idx, col_idx, item)
        self.csv_preview_table.resizeRowsToContents()

    def _dedupe_rows_by_manufacturer_data(self, rows):
        mfr_col = RESULT_HEADERS.index('MANUFACTURER_DATA')
        out = []
        index_by_mfr = {}
        for row in rows:
            normalized = list(row[:len(RESULT_HEADERS)]) + [''] * max(0, len(RESULT_HEADERS) - len(row))
            mfr = normalized[mfr_col]
            if mfr:
                old_idx = index_by_mfr.get(mfr)
                if old_idx is not None:
                    out.pop(old_idx)
                    index_by_mfr = {
                        existing[mfr_col]: idx
                        for idx, existing in enumerate(out)
                        if len(existing) > mfr_col and existing[mfr_col]
                    }
                index_by_mfr[mfr] = len(out)
            out.append(normalized[:len(RESULT_HEADERS)])
        return out

    def _export_csv_to_excel(self):
        if OPENPYXL_IMPORT_ERROR is not None:
            raise RuntimeError(f'openpyxl 未安装或导入失败: {OPENPYXL_IMPORT_ERROR}')
        rows = self._dedupe_rows_by_manufacturer_data(self._read_csv_rows())
        self._write_csv_rows(rows)
        self._refresh_csv_preview(rows)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Factory Test'
        ws.append(RESULT_HEADERS)
        for row in rows:
            ws.append(row)

        header_fill = PatternFill('solid', fgColor='D9EAF7')
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col_cells in ws.columns:
            column_letter = col_cells[0].column_letter
            max_len = max(len('' if cell.value is None else str(cell.value)) for cell in col_cells)
            ws.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 48)

        os.makedirs(os.path.dirname(os.path.abspath(self.result_xlsx_path)), exist_ok=True)
        wb.save(self.result_xlsx_path)

    def _on_confirm_clicked(self):
        dev = self._current_device()
        if dev is None:
            self._append_log('未连接设备，无法保存测试结果')
            return
        if not self.factory_ready:
            self._append_log('请先进入测试状态')
            return
        if not all(self.module_results[k] in ('成功', '失败') for k in ('radar', 'motor', 'laser')):
            self._append_log('三个模块都需要先选择成功/失败')
            return
        try:
            self._save_csv_record(dev)
        except Exception as ex:
            self._append_log(f'保存 CSV 失败: {ex}')
            return
        self._append_log(f'测试结果已保存到 CSV: {self.result_csv_path}')
        self._reset_module_results()
        self._update_buttons()

    def _on_export_excel_clicked(self):
        try:
            self._export_csv_to_excel()
        except Exception as ex:
            self._append_log(f'CSV 转 Excel 失败: {ex}')
            return
        self._append_log(f'CSV 已转换为 Excel: {self.result_xlsx_path}')

    def _clear_logs(self):
        self.host_log_text.clear()
        self.firmware_log_text.clear()

    def _append_log(self, line: str):
        now = _dt.datetime.now().strftime('%H:%M:%S.%f')[:-3]
        self.host_log_text.appendPlainText(f'[{now}] {line}')

    def _append_firmware_log(self, line: str):
        now = _dt.datetime.now().strftime('%H:%M:%S.%f')[:-3]
        self.firmware_log_text.appendPlainText(f'[{now}] {line}')

    def closeEvent(self, event):
        self.worker.stop()
        super().closeEvent(event)



def main():
    app = QtWidgets.QApplication(sys.argv)
    win = FactoryTestWindow()
    win.show()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()

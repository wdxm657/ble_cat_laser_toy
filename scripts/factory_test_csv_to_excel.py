# -*- coding: utf-8 -*-
"""
Convert factory_test_results.csv to an Excel workbook.

Dependencies:
    pip install openpyxl

Usage:
    python factory_test_csv_to_excel.py
    python factory_test_csv_to_excel.py input.csv output.xlsx
"""

import argparse
import csv
import os
import sys
from typing import List

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except Exception as ex:  # pragma: no cover - shown at runtime
    Workbook = None
    OPENPYXL_IMPORT_ERROR = ex
else:
    OPENPYXL_IMPORT_ERROR = None


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CSV = os.path.join(SCRIPT_DIR, "factory_test_results.csv")
DEFAULT_XLSX = os.path.join(SCRIPT_DIR, "factory_test_results.xlsx")


def _read_csv(path: str) -> List[List[str]]:
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            with open(path, "r", newline="", encoding=encoding) as fp:
                return list(csv.reader(fp))
        except UnicodeDecodeError:
            continue
    with open(path, "r", newline="") as fp:
        return list(csv.reader(fp))


def convert_csv_to_excel(csv_path: str, xlsx_path: str) -> None:
    if OPENPYXL_IMPORT_ERROR is not None:
        raise RuntimeError(f"openpyxl 未安装或导入失败: {OPENPYXL_IMPORT_ERROR}")
    if not os.path.exists(csv_path):
        raise FileNotFoundError(f"CSV 文件不存在: {csv_path}")

    rows = _read_csv(csv_path)
    if not rows:
        raise ValueError(f"CSV 文件为空: {csv_path}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Factory Test"

    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_cells in ws.columns:
        max_len = 0
        column_letter = col_cells[0].column_letter
        for cell in col_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 48)

    os.makedirs(os.path.dirname(os.path.abspath(xlsx_path)) or ".", exist_ok=True)
    wb.save(xlsx_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert W2M factory test CSV to xlsx.")
    parser.add_argument("csv", nargs="?", default=DEFAULT_CSV, help="input CSV path")
    parser.add_argument("xlsx", nargs="?", default=DEFAULT_XLSX, help="output xlsx path")
    args = parser.parse_args()

    try:
        convert_csv_to_excel(args.csv, args.xlsx)
    except Exception as ex:
        print(f"转换失败: {ex}", file=sys.stderr)
        return 1

    print(f"转换完成: {args.xlsx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
